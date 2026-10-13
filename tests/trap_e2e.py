# -*- coding: utf-8 -*-
"""第四位挑刺官的 8 章陷阱文档 + 本地后端全链路 E2E：登录→上传docx→解析→LLM抽题"""
import json, time, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import requests
from docx import Document

BASE = 'http://127.0.0.1:8000'
OUT = r'D:/query/tests/trap_docx_result.json'

# ---------- 1. 造 8 章陷阱 docx ----------
doc = Document()
doc.add_heading('网络与基础概念复习资料（挑刺专用测试版）', level=1)

doc.add_heading('第一章 普通知识点', level=2)
doc.add_paragraph('HTTP 协议本身是无状态的，服务器不会记住上一次请求是谁发的；'
                  'Cookie 存在客户端，Session 存在服务端，两者配合维持会话状态。')
doc.add_paragraph('栈是后进先出（LIFO）的数据结构，队列是先进先出（FIFO）的数据结构。'
                  '二分查找的时间复杂度是 O(log n)，前提是数组已经有序。')

doc.add_heading('第二章 易混淆概念对', level=2)
doc.add_paragraph('TCP 面向连接、提供可靠传输、有拥塞控制和重传机制，头部最少 20 字节；'
                  'UDP 无连接、不保证可靠、头部固定 8 字节，适合直播、语音、游戏等低延迟场景。')
doc.add_paragraph('功率的单位是瓦特（W），表示能量转换的快慢；电能的单位是瓦时（Wh），表示总共用了多少能量。'
                  '1 度电 = 1 千瓦时（kWh）= 3.6×10^6 焦耳。')

doc.add_heading('第三章 数字与日期密集段落', level=2)
doc.add_paragraph('IPv4 地址长度是 32 位，IPv6 地址长度是 128 位。')
doc.add_paragraph('以太网传统 MTU 是 1500 字节。')
doc.add_paragraph('2011 年 6 月 8 日是"世界 IPv6 日"测试；2019 年 11 月 25 日，RIPE NCC 宣布 IPv4 地址池正式耗尽。')

doc.add_heading('第四章 TCP 连接建立与关闭状态表', level=2)
tbl = doc.add_table(rows=1, cols=3)
hdr = tbl.rows[0].cells
hdr[0].text = '当前状态'
hdr[1].text = '触发事件'
hdr[2].text = '下一状态'
rows = [
    ('LISTEN', '收到 SYN', 'SYN_RCVD'),
    ('SYN_RCVD', '收到 ACK', 'ESTABLISHED'),
    ('ESTABLISHED', '主动关闭方发送 FIN', 'FIN_WAIT_1'),
    ('ESTABLISHED', '被动关闭方收到 FIN', 'CLOSE_WAIT'),
    ('CLOSE_WAIT', '应用发送 FIN', 'LAST_ACK'),
]
for r_ in rows:
    c = tbl.add_row().cells
    c[0].text, c[1].text, c[2].text = r_

doc.add_heading('第五章 代码片段', level=2)
doc.add_paragraph('def f(n):\n    s = 0\n    for i in range(1, n+1):\n        s += i\n    return s\nprint(f(4))')

doc.add_heading('第六章 前端术语（中英混排）', level=2)
doc.add_paragraph('React 使用 Virtual DOM（虚拟文档对象模型）来减少对真实 DOM 的直接操作。'
                  'useEffect 用于处理副作用（side effects），例如数据请求和订阅。'
                  'State（状态）更新会触发组件的重新渲染（re-render）。')

doc.add_heading('第七章 阅读判断（原文可能有误）', level=2)
doc.add_paragraph('HTTP/3 的传输层使用 TCP 协议，所以它比 HTTP/2 更可靠。')
doc.add_paragraph('Python 2 是目前的主流版本，新项目应该优先使用 Python 2。')

doc.add_heading('第八章 随笔', level=2)
doc.add_paragraph('学习是一件需要坚持的事情，只要坚持下去就一定会有收获。加油，明天会更好！')

DOCX = r'D:/query/tests/trap_test.docx'
doc.save(DOCX)
print('[1] 陷阱文档已生成:', DOCX)

# ---------- 2. 登录（本地后端 keep-alive 有怪癖：每次独立连接） ----------
def post(url, **kw):
    kw.setdefault('timeout', 20)
    kw.setdefault('headers', {})
    kw['headers'] = {**kw['headers'], 'Connection': 'close'}
    return requests.post(url, **kw)

def get(url, **kw):
    kw.setdefault('timeout', 20)
    kw.setdefault('headers', {})
    kw['headers'] = {**kw['headers'], 'Connection': 'close'}
    return requests.get(url, **kw)

r = post(BASE + '/api/auth/send-code', json={'phone': '13800001234'}).json()
code = r.get('demo_code')
assert code, r
time.sleep(1.5)
r = post(BASE + '/api/auth/login', json={'phone': '13800001234', 'code': code}).json()
assert r.get('token'), r
tok = r['token']
H = {'Authorization': 'Bearer ' + tok, 'Connection': 'close'}
print('[2] 登录 OK, 用户:', r['user']['name'])

# ---------- 3. 上传 docx ----------
with open(DOCX, 'rb') as f:
    up = post(BASE + '/api/doc/upload', files={'file': ('trap_test.docx', f)},
              headers=H, timeout=120).json()
assert up.get('ok'), up
jid = up['job_id']
print('[3] 上传 OK, job:', jid)
for _ in range(90):
    st = get(BASE + '/api/job/' + jid, headers=H, timeout=15).json()
    if st.get('status') == 'done':
        break
    time.sleep(1)
text = st.get('text') or ''
print('[4] 解析 OK, 正文字数:', len(text))
print('    正文前 200 字:', text[:200].replace('\n', ' | '))

# ---------- 4. LLM 异步抽题（按版块，单次调用小更稳） ----------
import re as _re
secs = []
cur = {'title': '开篇', 'lines': []}
for ln in text.splitlines():
    m = _re.match(r'^##\s*(.+)$', ln.strip())
    if m:
        if any(l.strip() for l in cur['lines']):
            secs.append({'title': cur['title'], 'text': '\n'.join(cur['lines'])})
        cur = {'title': m.group(1).strip(), 'lines': []}
    else:
        cur['lines'].append(ln)
if any(l.strip() for l in cur['lines']):
    secs.append({'title': cur['title'], 'text': '\n'.join(cur['lines'])})
print('   版块数:', len(secs), [s['title'] for s in secs])

body = {'text': '', 'sections': secs, 'maxPerSection': 6, 'qtype': 'auto',
        'apiKey': '', 'baseUrl': '', 'model': ''}
r = post(BASE + '/api/extract-job', json=body, headers=H, timeout=20).json()
assert r.get('ok'), r
xid = r['job_id']
print('[5] LLM 抽题任务已提交:', xid)
t0 = time.time()
while time.time() - t0 < 300:
    st = get(BASE + '/api/extract-job/' + xid, headers=H, timeout=15).json()
    if st.get('status') == 'done':
        break
    if st.get('status') == 'error':
        print('抽取失败:', st); sys.exit(1)
    time.sleep(3)
items = st.get('items') or []
print('[6] 抽题完成, 用时 %.1fs, 共 %d 题' % (time.time() - t0, len(items)))

json.dump({'text': text, 'items': items}, open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print('结果已存:', OUT)
for i, q in enumerate(items, 1):
    print('---Q%d [%s] topic=%s' % (i, q.get('qtype'), q.get('topic', '')))
    print('  问:', q.get('question'))
    ans = (q.get('answer') or '')[:160]
    print('  答:', ans)
    if q.get('options'):
        print('  选项:', ' | '.join(q.get('options')), '正确:', q.get('correct'))
